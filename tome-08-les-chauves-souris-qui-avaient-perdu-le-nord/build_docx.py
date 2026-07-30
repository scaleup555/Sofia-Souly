#!/usr/bin/env python3
"""Build the final .docx for Tome 8 — Les Chauves-souris qui avaient perdu le nord."""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt
from docx.oxml.ns import qn

BASE = Path(__file__).parent
CHAPTERS_DIR = BASE / "chapitres"
OUTPUT = BASE / "Sofia-et-Souly-Tome-8-Les-Chauves-souris-qui-avaient-perdu-le-nord.docx"

TITLE = "Les Chauves-souris qui avaient perdu le nord"
SUBTITLE = "Une aventure en chiroptérologie"
TOME_LABEL = "Tome 8"
JURY_SCORE = "8,8 / 10"

CHAPTER_FILES = [f"chapitre-{i:02d}.md" for i in range(1, 13)]

FONT_NAME = "Georgia"


def set_default_font(document):
    style = document.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = Pt(11.5)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), FONT_NAME)


def set_page_geometry(section):
    section.page_width = Inches(5.5)
    section.page_height = Inches(8.5)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.55)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)


def add_centered(document, text, size, bold=False, italic=False, space_after=None):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    if text:
        run = p.add_run(text)
        run.font.name = FONT_NAME
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
    return p


def add_blank(document, count=1):
    for _ in range(count):
        document.add_paragraph()


def new_page_section(document):
    section = document.add_section(WD_SECTION.NEW_PAGE)
    set_page_geometry(section)
    return section


def parse_chapter(path):
    text = path.read_text(encoding="utf-8")
    lines = text.splitlines()
    header = lines[0]
    m = re.match(r"#\s*Chapitre\s+(\d+)\s*—\s*(.+)", header)
    number, title = m.group(1), m.group(2).strip()
    body = "\n".join(lines[1:]).strip("\n")
    blocks = [b.strip() for b in re.split(r"\n\s*\n", body) if b.strip()]
    return number, title, blocks


def add_body_block(document, block):
    if block.strip() == "---":
        p = add_centered(document, "• • •", 11)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(10)
        return
    paragraph_text = " ".join(line.strip() for line in block.splitlines())
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(paragraph_text)
    run.font.name = FONT_NAME
    run.font.size = Pt(11.5)


def build():
    document = Document()
    set_default_font(document)
    set_page_geometry(document.sections[0])

    # --- Title page ---
    add_blank(document, 4)
    add_centered(document, "Sofia & Souly", 30, bold=True)
    add_centered(document, TOME_LABEL, 14, italic=True)
    add_blank(document, 1)
    add_centered(document, TITLE, 20, bold=True)
    add_blank(document, 1)
    add_centered(document, SUBTITLE, 11, italic=True)
    add_blank(document, 10)
    add_centered(document, "Roman jeunesse — 8 à 14 ans", 9.5)
    add_blank(document, 7)
    add_centered(document, f"Sofia & Souly — {TOME_LABEL} : {TITLE}", 9.5)
    add_blank(document, 1)
    add_centered(document, "Une histoire originale.", 9.5)
    add_centered(
        document,
        "Personnages et lieux fictifs, à l'exception des repères",
        9.5,
    )
    add_centered(document, "scientifiques réels mentionnés dans le récit.", 9.5)
    add_blank(document, 1)
    add_centered(
        document,
        "Édition numérique — manuscrit validé par un jury indépendant",
        9.5,
    )
    add_centered(document, f"de littérature jeunesse (note finale : {JURY_SCORE}).", 9.5)

    # --- Table of contents ---
    new_page_section(document)
    add_centered(document, "Table des matières", 16, bold=True)
    add_blank(document, 1)

    chapters = [parse_chapter(CHAPTERS_DIR / fname) for fname in CHAPTER_FILES]

    for number, title, _ in chapters:
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(f"Chapitre {number} — {title}")
        run.font.name = FONT_NAME
        run.font.size = Pt(11)

    # --- Chapters ---
    for number, title, blocks in chapters:
        new_page_section(document)
        add_blank(document, 2)
        add_centered(document, f"CHAPITRE {number}", 12, bold=True)
        add_centered(document, title, 17, bold=True, space_after=18)
        for block in blocks:
            add_body_block(document, block)

    document.save(OUTPUT)
    print(f"Saved {OUTPUT}")


if __name__ == "__main__":
    build()
